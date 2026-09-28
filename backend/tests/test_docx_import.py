from __future__ import annotations

from docx import Document
from docx.enum.section import WD_SECTION

from scripts.clean_standard_prose import clean_snippet_prose
from scripts.import_docx_standards import (
    _match_by_title,
    _render_table,
    extract_docx_snippets,
)


def test_docx_import_preserves_paragraphs_tables_and_sections(tmp_path) -> None:
    path = tmp_path / "GB50096-2011住宅设计规范.docx"
    document = Document()
    document.add_paragraph("1.0.1 住宅设计应符合本规范的规定，并应满足公共交")
    document.add_paragraph("通及相邻户门开启的要求。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "类别"
    table.cell(0, 1).text = "洞口宽度(m)"
    table.cell(1, 0).text = "共用外门"
    table.cell(1, 1).text = "1.20"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("2.0.1 住宅 residential buildings")
    document.save(path)

    snippets, stats = extract_docx_snippets(path)

    assert stats["pages"] == 2
    assert len(snippets) == 2
    assert snippets[0]["article"] == "1.0.1"
    assert "公共交通及相邻户门开启" in snippets[0]["text"]
    assert "| 类别 | 洞口宽度(m) |" in snippets[0]["text"]
    assert "| 共用外门 | 1.20 |" in snippets[0]["text"]
    assert snippets[1]["page"] == 2
    assert snippets[1]["text"] == "2.0.1 住宅 residential buildings"


def test_docx_import_preserves_merged_cell_spans() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "类别"
    table.cell(0, 1).text = "宽度"
    table.cell(0, 2).text = "高度"
    table.cell(1, 0).text = "住宅"
    table.cell(1, 1).text = "0.26"
    table.cell(1, 2).text = "0.175"
    table.cell(2, 0).text = ""
    table.cell(2, 1).text = "0.28"
    table.cell(2, 2).text = "0.165"
    table.cell(0, 2).text = ""
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(0, 1).merge(table.cell(0, 2))

    rendered = _render_table(table)

    assert "| 类别 | 宽度 @colspan=2 |" in rendered
    assert "| 住宅 @rowspan=2 | 0.26 | 0.175 |" in rendered
    assert "| 0.28 | 0.165 |" in rendered


def test_prose_cleanup_preserves_tables_and_repairs_spacing() -> None:
    text = (
        "1.60 x\n"
        "0.60 = 0.96 (m Z )。\n"
        "| 类别 | 数值 |\n"
        "| 办公 | 0.96 |"
    )

    cleaned = clean_snippet_prose(text)

    assert "1.60×0.60=0.96(m²)。" in cleaned
    assert "\n| 类别 | 数值 |\n" in cleaned
    assert cleaned.endswith("| 办公 | 0.96 |")


def test_title_match_does_not_cross_region(tmp_path) -> None:
    source = (
        tmp_path
        / "\u6d59\u6c5f\u7701"
        / "\u676d\u5dde\u5e02\u57ce\u5e02\u89c4\u5212\u7ba1\u7406\u6280\u672f\u89c4\u5b9a.docx"
    )
    source.parent.mkdir()
    source.write_bytes(b"docx")
    records = [
        (
            tmp_path / "chengdu.json",
            {},
            {
                "title": "\u6210\u90fd\u5e02\u57ce\u5e02\u89c4\u5212\u7ba1\u7406\u6280\u672f\u89c4\u5b9a2024",
                "region": "\u6210\u90fd\u5e02",
                "province": "\u56db\u5ddd\u7701",
                "city": "\u6210\u90fd\u5e02",
                "district": "",
            },
        )
    ]

    assert _match_by_title(source, records) is None
